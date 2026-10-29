from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "203748"
MUTED = "667085"
TABLE_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
BORDER = "B8C4D0"
BODY_FONT = "Calibri"
CJK_FONT = "Arial Unicode MS"
CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120

DEFAULT_DOCUMENT_META = {
    "running_header": "欢畅流量翻番 | 中移与北移订购接口",
    "kicker": "技术业务订购手册",
    "title": "欢畅流量翻番",
    "subtitle": "中移与北移业务订购流程及接口说明",
    "summary": "基于页面源码与 HAR 实例的字段级还原",
    "version": "V1.0",
    "analysis_date": "2026-07-24",
    "audience": "产品、前端、后端、测试、联调与运维人员",
    "evidence_scope": "中移明文还原；北移源码字段与抓包时序交叉验证",
    "core_title": "欢畅流量翻番：中移与北移业务订购流程及接口说明",
    "core_subject": "中移与北京移动业务订购页面、接口、字段和状态机分析",
    "core_keywords": "CMCC, 北京移动, 业务订购, HAR, 接口文档",
}


def set_run_font(run, size=None, bold=None, italic=None, color=None, name=BODY_FONT):
    actual_name = CJK_FONT if CJK_RE.search(run.text or "") else name
    run.font.name = actual_name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), actual_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), actual_name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, size, color="000000", bold=False):
    style.font.name = CJK_FONT
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), CJK_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), CJK_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    cant_split.set(qn("w:val"), "true")


def apply_table_geometry(table, widths):
    if sum(widths) != CONTENT_DXA:
        raise ValueError(f"Table widths must sum to {CONTENT_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths[index] / 1440)
            set_cell_margins(cell)


def add_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def table_widths(headers):
    count = len(headers)
    normalized = [re.sub(r"[`*]", "", h).strip() for h in headers]
    if count == 2:
        return [2500, 6860]
    if count == 3:
        return [2200, 1700, 5460]
    if count == 4:
        if any("必填" in h for h in normalized):
            return [1900, 1000, 1100, 5360]
        return [1800, 2100, 2500, 2960]
    if count == 5:
        if normalized[0] in {"序号", "用例"}:
            return [700, 2500, 1200, 2800, 2160]
        if "类型" in normalized and any("来源" in h for h in normalized):
            return [1700, 900, 1100, 2100, 3560]
        return [1500, 1700, 1600, 2200, 2360]
    return [CONTENT_DXA // count] * (count - 1) + [CONTENT_DXA - (CONTENT_DXA // count) * (count - 1)]


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])
    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, size=8.5, color=MUTED)


def create_numbering(document):
    numbering = document.part.numbering_part.element

    def next_id(tag, attr):
        values = [int(node.get(qn(attr))) for node in numbering.findall(qn(tag))]
        return max(values, default=0) + 1

    def add_definition(marker, num_fmt):
        abstract_id = next_id("w:abstractNum", "w:abstractNumId")
        num_id = next_id("w:num", "w:numId")
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_fmt)
        level.append(fmt)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), marker)
        level.append(text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        level.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "540")
        ind.set(qn("w:hanging"), "270")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "80")
        spacing.set(qn("w:line"), "300")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        level.append(p_pr)
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), BODY_FONT)
        fonts.set(qn("w:hAnsi"), BODY_FONT)
        fonts.set(qn("w:eastAsia"), CJK_FONT)
        r_pr.append(fonts)
        level.append(r_pr)
        abstract.append(level)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)
        return num_id

    return add_definition("●", "bullet"), add_definition("%1.", "decimal")


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])


INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
CJK_RE = re.compile(r"[\u3400-\u9fff\uf900-\ufaff]")


def add_inline(paragraph, text, size=None, color=None, bold=False):
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run, size=size, color=color, bold=bold)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Courier New", size=(size or 10.5) - 0.5, color=DARK_BLUE)
            run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=color, bold=bold)


def split_table_row(line):
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    cells = []
    current = []
    in_code = False
    escaped = False
    for char in text:
        if escaped:
            current.append(char)
            escaped = False
        elif char == "\\":
            escaped = True
            current.append(char)
        elif char == "`":
            in_code = not in_code
            current.append(char)
        elif char == "|" and not in_code:
            cells.append("".join(current).strip())
            current = []
        else:
            current.append(char)
    cells.append("".join(current).strip())
    return cells


def is_table_separator(line):
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_data_table(document, rows):
    headers = rows[0]
    widths = table_widths(headers)
    table = document.add_table(rows=len(rows), cols=len(headers))
    table.style = "Table Grid"
    apply_table_geometry(table, widths)
    add_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        set_row_cant_split(row)
    center_headers = {"序号", "类型", "必填", "必填级别", "级别", "本次响应", "来源"}
    for row_index, row_values in enumerate(rows):
        for col_index, value in enumerate(row_values):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_index == 0:
                set_cell_shading(cell, TABLE_FILL)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            header_text = re.sub(r"[`*]", "", headers[col_index])
            if row_index == 0 or header_text in center_headers:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(
                paragraph,
                value.replace("  ", " "),
                size=8.5,
                color="000000",
                bold=(row_index == 0),
            )
    after = document.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(4)
    return table


def add_callout(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.05)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.2
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_FILL)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    p_pr.append(borders)
    add_inline(paragraph, text, size=10, color="25313C")


def add_code_block(document, lines):
    paragraph = document.add_paragraph(style="Code Block")
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, name="Courier New", size=8.2, color="1F2937")
        run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)


def configure_document(document, metadata):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = document.styles["Normal"]
    set_style_font(normal, 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    h1 = document.styles["Heading 1"]
    set_style_font(h1, 16, BLUE, True)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = document.styles["Heading 2"]
    set_style_font(h2, 13, BLUE, True)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(7)
    h2.paragraph_format.keep_with_next = True

    h3 = document.styles["Heading 3"]
    set_style_font(h3, 12, DARK_BLUE, True)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(5)
    h3.paragraph_format.keep_with_next = True

    code_style = document.styles.add_style("Code Block", 1)
    set_style_font(code_style, 8.2, "1F2937")
    code_style.font.name = "Courier New"
    code_style.paragraph_format.left_indent = Inches(0.12)
    code_style.paragraph_format.right_indent = Inches(0.05)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(8)
    code_style.paragraph_format.line_spacing = 1.0
    code_p_pr = code_style._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F4F7")
    code_p_pr.append(shd)

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.paragraph_format.space_after = Pt(0)
    add_inline(header_p, metadata["running_header"], size=8.5, color=MUTED)

    footer_p = section.footer.paragraphs[0]
    footer_p.paragraph_format.space_after = Pt(0)
    add_page_number(footer_p)


def add_cover(document, metadata):
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(90)

    kicker = document.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    add_inline(kicker, metadata["kicker"], size=11, color=BLUE, bold=True)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    add_inline(title, metadata["title"], size=29, color=NAVY, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    add_inline(subtitle, metadata["subtitle"], size=16, color=DARK_BLUE)

    summary = document.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.space_after = Pt(70)
    add_inline(summary, metadata["summary"], size=10.5, color=MUTED)

    metadata = [
        ("版本", metadata["version"]),
        ("分析日期", metadata["analysis_date"]),
        ("适用对象", metadata["audience"]),
        ("证据范围", metadata["evidence_scope"]),
    ]
    for label, value in metadata:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        add_inline(p, f"{label}：", size=9.5, color=MUTED, bold=True)
        add_inline(p, value, size=9.5, color="334155")

    document.add_page_break()


def add_contents(document, lines):
    heading = document.add_paragraph("目录", style="Heading 1")
    heading.paragraph_format.space_before = Pt(0)
    for line in lines:
        text = line[3:].strip()
        p = document.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.16)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = False
        add_inline(p, text, size=10, color=DARK_BLUE)
    document.add_page_break()


def build(markdown_path, output_path, metadata=None):
    resolved_metadata = {**DEFAULT_DOCUMENT_META, **(metadata or {})}
    lines = Path(markdown_path).read_text(encoding="utf8").splitlines()
    document = Document()
    configure_document(document, resolved_metadata)
    bullet_num_id, decimal_num_id = create_numbering(document)
    add_cover(document, resolved_metadata)
    add_contents(document, [line for line in lines if line.startswith("## ")])

    start = next(i for i, line in enumerate(lines) if line.startswith("## "))
    index = start
    in_numbered_list = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            in_numbered_list = False
            index += 1
            continue

        if stripped.startswith("```"):
            in_numbered_list = False
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(document, code_lines)
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            in_numbered_list = False
            rows = [split_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_table_row(lines[index]))
                index += 1
            expected = len(rows[0])
            rows = [row[:expected] + [""] * max(0, expected - len(row)) for row in rows]
            add_data_table(document, rows)
            continue

        if stripped.startswith("### "):
            p = document.add_paragraph(style="Heading 2")
            add_inline(p, stripped[4:].strip())
        elif stripped.startswith("## "):
            p = document.add_paragraph(style="Heading 1")
            add_inline(p, stripped[3:].strip())
        elif stripped.startswith("> "):
            add_callout(document, stripped[2:].strip())
        elif stripped.startswith("- "):
            p = document.add_paragraph()
            apply_numbering(p, bullet_num_id)
            add_inline(p, stripped[2:].strip())
        elif re.match(r"^\d+\.\s+", stripped):
            if not in_numbered_list:
                _, decimal_num_id = create_numbering(document)
            p = document.add_paragraph()
            apply_numbering(p, decimal_num_id)
            add_inline(p, re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped.startswith("**") and stripped.endswith("**"):
            p = document.add_paragraph(style="Heading 3")
            add_inline(p, stripped[2:-2].strip())
        else:
            p = document.add_paragraph()
            add_inline(p, stripped.replace("  ", " "))
        in_numbered_list = bool(re.match(r"^\d+\.\s+", stripped))
        index += 1

    core = document.core_properties
    core.title = resolved_metadata["core_title"]
    core.subject = resolved_metadata["core_subject"]
    core.author = ""
    core.last_modified_by = ""
    core.keywords = resolved_metadata["core_keywords"]
    document.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) not in {3, 4}:
        raise SystemExit(
            "Usage: _build_order_document.py <input.md> <output.docx> [metadata.json]"
        )
    custom_metadata = None
    if len(sys.argv) == 4:
        custom_metadata = json.loads(Path(sys.argv[3]).read_text(encoding="utf8"))
    build(sys.argv[1], sys.argv[2], custom_metadata)
