from __future__ import annotations

import sys
import tempfile
import uuid
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

import _build_order_document as base


DOCUMENTS = [
    {
        "source": "统一业务订购系统-开发规范.md",
        "output": "统一业务订购系统-开发规范.docx",
        "running_header": "统一业务订购系统 | 开发规范",
        "kicker": "工程与安全规范",
        "title": "统一业务订购系统",
        "subtitle": "开发规范",
        "summary": "北移、中移商城与网格通统一快捷订购工程基线",
        "audience": "开发、测试、安全、联调、运维与技术负责人",
        "evidence_scope": "基于现有系统、保存源码和三平台 HAR 分析",
        "core_title": "统一业务订购系统开发规范",
        "core_subject": "三平台业务订购系统的工程、安全、接口与测试规范",
        "core_keywords": "CMCC, 业务订购, 开发规范, 安全, 幂等",
    },
    {
        "source": "统一业务订购系统-技术说明.md",
        "output": "统一业务订购系统-技术说明.docx",
        "running_header": "统一业务订购系统 | 技术说明",
        "kicker": "协议与接口参考",
        "title": "统一业务订购系统",
        "subtitle": "技术说明",
        "summary": "三平台认证、加密、接口、字段和适配器契约",
        "audience": "后端、前端、测试、联调、安全与运维人员",
        "evidence_scope": "HAR 明文还原、页面源码与现有项目实现交叉验证",
        "core_title": "统一业务订购系统技术说明",
        "core_subject": "北移、中移商城和网格通订购协议及统一 API",
        "core_keywords": "CMCC, 北京移动, 中移商城, 网格通, API",
    },
    {
        "source": "统一业务订购系统-系统设计.md",
        "output": "统一业务订购系统-系统设计.docx",
        "running_header": "统一业务订购系统 | 系统设计",
        "kicker": "架构与交互设计",
        "title": "统一业务订购系统",
        "subtitle": "系统设计",
        "summary": "统一工作台、状态机、模块、数据与部署设计",
        "audience": "产品、架构、开发、测试、联调与运维人员",
        "evidence_scope": "面向 bj-mobile-business-local 的可实施设计",
        "core_title": "统一业务订购系统设计文档",
        "core_subject": "统一快捷订购工作台的架构、交互和部署设计",
        "core_keywords": "CMCC, 系统设计, 状态机, Vue, Express",
    },
]

EMBEDDED_FONT_NAME = "Arial Unicode MS"
EMBEDDED_FONT_PATH = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")


def _obfuscate_font(font_data: bytes, key: uuid.UUID) -> bytes:
    output = bytearray(font_data)
    guid_bytes = key.bytes
    for index in range(min(32, len(output))):
        output[index] ^= guid_bytes[15 - (index % 16)]
    return bytes(output)


def embed_font(docx_path: Path, font_path: Path, font_name: str):
    if not font_path.exists():
        raise FileNotFoundError(font_path)
    relationship_id = "rIdEmbeddedCjk"
    key = uuid.uuid4()
    target = "fonts/embedded-cjk.odttf"
    with ZipFile(docx_path, "r") as source:
        entries = {name: source.read(name) for name in source.namelist()}

    font_table = parse_xml(entries["word/fontTable.xml"])
    font_node = None
    for candidate in font_table.findall(qn("w:font")):
        if candidate.get(qn("w:name")) == font_name:
            font_node = candidate
            break
    if font_node is None:
        font_node = OxmlElement("w:font")
        font_node.set(qn("w:name"), font_name)
        font_table.append(font_node)
    for existing in list(font_node.findall(qn("w:embedRegular"))):
        font_node.remove(existing)
    embedded = OxmlElement("w:embedRegular")
    embedded.set(qn("r:id"), relationship_id)
    embedded.set(qn("w:fontKey"), "{" + str(key).upper() + "}")
    font_node.append(embedded)
    entries["word/fontTable.xml"] = etree.tostring(
        font_table,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )

    rels_path = "word/_rels/fontTable.xml.rels"
    rels_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    if rels_path in entries:
        rels = parse_xml(entries[rels_path])
    else:
        rels = parse_xml(f'<Relationships xmlns="{rels_namespace}"/>')
    for relation in list(rels):
        if relation.get("Id") == relationship_id:
            rels.remove(relation)
    relation = etree.Element(f"{{{rels_namespace}}}Relationship")
    relation.set("Id", relationship_id)
    relation.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font")
    relation.set("Target", target)
    rels.append(relation)
    entries[rels_path] = etree.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)

    content_types = parse_xml(entries["[Content_Types].xml"])
    content_namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    if not any(node.get("Extension") == "odttf" for node in content_types):
        default = etree.Element(f"{{{content_namespace}}}Default")
        default.set("Extension", "odttf")
        default.set("ContentType", "application/vnd.openxmlformats-officedocument.obfuscatedFont")
        content_types.append(default)
    entries["[Content_Types].xml"] = etree.tostring(
        content_types,
        xml_declaration=True,
        encoding="UTF-8",
        standalone=True,
    )
    entries[f"word/{target}"] = _obfuscate_font(font_path.read_bytes(), key)

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, dir=docx_path.parent) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with ZipFile(temporary_path, "w", ZIP_DEFLATED) as output:
            for name, data in entries.items():
                output.writestr(name, data)
        temporary_path.replace(docx_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def add_memo_masthead(document, metadata):
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(14)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    base.add_inline(kicker, metadata["kicker"].upper(), size=10, color=base.BLUE, bold=True)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    base.add_inline(title, metadata["title"], size=24, color="000000", bold=True)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    base.add_inline(subtitle, metadata["subtitle"], size=15, color="373737", bold=True)

    summary = document.add_paragraph()
    summary.paragraph_format.space_after = Pt(18)
    summary.paragraph_format.line_spacing = 1.2
    base.add_inline(summary, metadata["summary"], size=10.5, color=base.MUTED)

    rows = [
        ("版本", metadata["version"]),
        ("日期", metadata["analysis_date"]),
        ("适用对象", metadata["audience"]),
        ("依据", metadata["evidence_scope"]),
        ("状态", "开发与验收基线"),
    ]
    for label, value in rows:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.15
        base.add_inline(paragraph, f"{label}：", size=10.5, color="000000", bold=True)
        base.add_inline(paragraph, value, size=10.5, color="000000")

    rule = document.add_paragraph()
    rule.paragraph_format.space_before = Pt(13)
    rule.paragraph_format.space_after = Pt(0)
    rule.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_pr = rule._p.get_or_add_pPr()
    borders = base.OxmlElement("w:pBdr")
    bottom = base.OxmlElement("w:bottom")
    bottom.set(base.qn("w:val"), "single")
    bottom.set(base.qn("w:sz"), "8")
    bottom.set(base.qn("w:space"), "1")
    bottom.set(base.qn("w:color"), base.BLUE)
    borders.append(bottom)
    p_pr.append(borders)
    document.add_page_break()


def main(output_dir: Path):
    output_dir.mkdir(parents=True, exist_ok=True)
    base.CJK_FONT = EMBEDDED_FONT_NAME
    base.add_cover = add_memo_masthead
    source_dir = Path(__file__).resolve().parent
    for config in DOCUMENTS:
        metadata = {
            **config,
            "version": "V1.0",
            "analysis_date": "2026-07-24",
        }
        output_path = output_dir / config["output"]
        base.build(
            source_dir / config["source"],
            output_path,
            metadata,
        )
        embed_font(output_path, EMBEDDED_FONT_PATH, EMBEDDED_FONT_NAME)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: _build_system_documents.py <output-dir>")
    main(Path(sys.argv[1]))
