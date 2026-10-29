from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont

import _build_order_document as base
from _build_system_documents import add_memo_masthead, embed_font


SOURCE_NAME = "中国移动云应用-HAR流程解析与实现说明.md"
OUTPUT_NAME = "中国移动云应用-HAR流程解析与实现说明.docx"

METADATA = {
    "running_header": "中国移动云应用 | HAR 流程解析与实现说明",
    "kicker": "协议、架构、接口与串流参考",
    "title": "中国移动云应用",
    "subtitle": "HAR 流程解析、系统设计与接口说明",
    "summary": "从离线证据还原到 BFF、实例编排、信令票据和安全部署的完整参考方案",
    "version": "V2.0",
    "analysis_date": "2026-08-11",
    "audience": "前端、服务端、联调、测试与安全人员",
    "evidence_scope": "两份 HAR、两份离线响应资源包与同版本前端 bundle 交叉验证",
    "core_title": "中国移动云应用 HAR 流程解析、系统设计与接口说明",
    "core_subject": "10086 单点登录、139 云应用门户、能力授权、实例编排、控制面起流、Socket.IO 与 WebRTC",
    "core_keywords": "CMCC, 10086, 139云应用, HAR, SSO, Cloudapp, BFF, Socket.IO, WebRTC",
}

FONT_PATH = Path("/System/Library/Fonts/STHeiti Medium.ttc")
EMBEDDED_CJK_FONT_PATH = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")


def _font(size: int, bold: bool = False):
    # STHeiti is used for both regular and emphasized labels to keep Chinese glyphs stable.
    return ImageFont.truetype(str(FONT_PATH), size=size)


def _color(value: str):
    return value if value.startswith("#") else f"#{value}"


def _rounded_box(draw, xy, fill, outline, radius=18, width=3):
    draw.rounded_rectangle(
        xy,
        radius=radius,
        fill=_color(fill),
        outline=_color(outline),
        width=width,
    )


def _center_text(draw, xy, text, font, fill="18212B"):
    x1, y1, x2, y2 = xy
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=7, align="center")
    width = bbox[2] - bbox[0]
    height = bbox[3] - bbox[1]
    draw.multiline_text(
        ((x1 + x2 - width) / 2, (y1 + y2 - height) / 2),
        text,
        font=font,
        fill=_color(fill),
        spacing=7,
        align="center",
    )


def _arrow(draw, start, end, color="2E74B5", width=6):
    draw.line([start, end], fill=_color(color), width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - direction * 18, ey - 11), (ex - direction * 18, ey + 11)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 11, ey - direction * 18), (ex + 11, ey - direction * 18)]
    draw.polygon(points, fill=_color(color))


def build_flow_figure(output_path: Path):
    image = Image.new("RGB", (1800, 1020), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = _font(43)
    lane_font = _font(27)
    body_font = _font(25)
    small_font = _font(22)

    draw.text((70, 36), "入口、会话、能力授权与串流主链路", font=title_font, fill="#1F4D78")
    draw.text((70, 94), "10086 跳转为可选入口；能否起流最终由 canPlay 与能力授权结果共同决定", font=small_font, fill="#59636E")

    lanes = [
        ("浏览器", 70, 330),
        ("10086 入口", 390, 650),
        ("139 门户", 710, 970),
        ("业务 API", 1030, 1290),
        ("控制面 / 信令", 1350, 1730),
    ]
    for label, x1, x2 in lanes:
        _rounded_box(draw, (x1, 150, x2, 212), "E8EEF5", "B8C4D0", radius=12, width=2)
        _center_text(draw, (x1, 150, x2, 212), label, lane_font, "1F4D78")
        draw.line([(x1 + x2) / 2, 222, (x1 + x2) / 2, 960], fill="#D6DCE3", width=2)

    _rounded_box(draw, (82, 260, 318, 340), "F4F6F9", "B8C4D0")
    _center_text(draw, (82, 260, 318, 340), "打开入口链接", body_font)
    _arrow(draw, (318, 300), (402, 300))
    _rounded_box(draw, (402, 250, 638, 350), "FFF8E5", "7A5A00")
    _center_text(draw, (402, 250, 638, 350), "bindAccount/new\n校验并 302", body_font)
    _arrow(draw, (638, 300), (722, 300))
    _rounded_box(draw, (722, 250, 958, 350), "EAF3FA", "2E74B5")
    _center_text(draw, (722, 250, 958, 350), "读取 Fragment\n解密 i / 清理 token", body_font)

    _arrow(draw, (958, 395), (1042, 395))
    _rounded_box(draw, (1042, 350, 1278, 440), "F4F6F9", "B8C4D0")
    _center_text(draw, (1042, 350, 1278, 440), "query / queryBrief\n获取应用配置", body_font)
    _arrow(draw, (1042, 485), (958, 485))
    _rounded_box(draw, (722, 445, 958, 525), "EAF3FA", "2E74B5")
    _center_text(draw, (722, 445, 958, 525), "可选：ssoHW\n入口 Token 换会话", body_font)

    _arrow(draw, (958, 575), (1042, 575))
    _rounded_box(draw, (1042, 530, 1278, 620), "F4F6F9", "B8C4D0")
    _center_text(draw, (1042, 530, 1278, 620), "四项并发准备\n授权 / RCS / 地域 / AKSK", body_font)

    _arrow(draw, (1278, 575), (1362, 575))
    _rounded_box(draw, (1362, 520, 1718, 630), "FFF0F0", "9B1C1C")
    _center_text(draw, (1362, 520, 1718, 630), "判定\ncanPlay=1 且能力授权成功？", body_font)

    _arrow(draw, (1540, 630), (1540, 705), color="2E74B5")
    _rounded_box(draw, (1362, 705, 1718, 790), "EAF3FA", "2E74B5")
    _center_text(draw, (1362, 705, 1718, 790), "是：autostreaming\n轮询 taskId 至 READY", body_font)
    _arrow(draw, (1540, 790), (1540, 865))
    _rounded_box(draw, (1362, 865, 1718, 950), "E8F4EC", "3C7A4F")
    _center_text(draw, (1362, 865, 1718, 950), "Socket.IO 建连\n创建播放器并进入 PLAYING", body_font)

    _arrow(draw, (1362, 575), (1265, 705), color="9B1C1C")
    _rounded_box(draw, (1030, 705, 1280, 790), "FFF0F0", "9B1C1C")
    _center_text(draw, (1030, 705, 1280, 790), "否：离线/授权异常\n不创建云实例", body_font, "7E1717")

    image.save(output_path, format="PNG", optimize=True)


def build_architecture_figure(output_path: Path):
    image = Image.new("RGB", (1800, 1120), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = _font(43)
    subtitle_font = _font(22)
    lane_font = _font(27)
    box_font = _font(25)
    note_font = _font(20)

    draw.text((70, 36), "推荐分层架构与信任边界", font=title_font, fill="#1F4D78")
    draw.text(
        (70, 94),
        "浏览器只持有短期引用与一次性票据；会话、签名、幂等和实例回收集中在服务端",
        font=subtitle_font,
        fill="#59636E",
    )

    lanes = [
        ("浏览器 / 不可信终端", 150, 315, "F4F6F9", "B8C4D0"),
        ("接入与编排平台", 365, 715, "EAF3FA", "2E74B5"),
        ("外部系统", 780, 980, "FFF8E5", "7A5A00"),
    ]
    for label, y1, y2, fill, outline in lanes:
        _rounded_box(draw, (70, y1, 1730, y2), fill, outline, radius=14, width=2)
        draw.text((92, y1 + 20), label, font=lane_font, fill=_color(outline))

    columns = [(235, 485), (600, 850), (965, 1215), (1330, 1580)]
    browser = ["入口消费 UI", "状态与错误渲染", "PlayerAdapter", "关闭 / 恢复控制"]
    for (x1, x2), label in zip(columns, browser):
        _rounded_box(draw, (x1, 220, x2, 285), "FFFFFF", "B8C4D0", radius=10, width=2)
        _center_text(draw, (x1, 220, x2, 285), label, box_font)

    core = ["EntryResolver", "SessionBroker", "LaunchOrchestrator", "TicketIssuer"]
    for (x1, x2), label in zip(columns, core):
        _rounded_box(draw, (x1, 435, x2, 510), "FFFFFF", "2E74B5", radius=10, width=2)
        _center_text(draw, (x1, 435, x2, 510), label, box_font, "1F4D78")

    shared = ["PolicyEngine", "CredentialBroker\n+ KMS", "Redis / 幂等状态", "TelemetryRedactor"]
    for (x1, x2), label in zip(columns, shared):
        _rounded_box(draw, (x1, 595, x2, 675), "F7FAFD", "7E9AB5", radius=10, width=2)
        _center_text(draw, (x1, 595, x2, 675), label, note_font, "334E68")

    external = ["10086 入口", "139 门户 API", "云应用控制面", "Socket.IO\n+ WebRTC"]
    for (x1, x2), label in zip(columns, external):
        _rounded_box(draw, (x1, 865, x2, 940), "FFFFFF", "7A5A00", radius=10, width=2)
        _center_text(draw, (x1, 865, x2, 940), label, box_font, "5C4700")

    for x1, x2 in columns:
        center = (x1 + x2) / 2
        _arrow(draw, (center, 285), (center, 435), color="2E74B5", width=4)
    _arrow(draw, (1710, 715), (1710, 780), color="7A5A00", width=4)
    draw.text((1510, 730), "受控上游调用", font=note_font, fill="#7A5A00")

    draw.text(
        (92, 1025),
        "所有跨边界调用携带内部 requestId / launchId；Token、AK/SK、手机号、完整 URL 与 SDP/ICE 不进入普通日志。",
        font=note_font,
        fill="#59636E",
    )
    image.save(output_path, format="PNG", optimize=True)


def add_code_or_flow(document, lines):
    if lines and lines[0].strip() == "sequenceDiagram":
        figure_path = Path(tempfile.gettempdir()) / "cmcc-cloudapp-main-flow.png"
        build_flow_figure(figure_path)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)
        picture = paragraph.add_run().add_picture(str(figure_path), width=Inches(6.3))
        picture._inline.docPr.set(
            "descr",
            "10086 可选入口、139 门户会话、并发能力准备、起流判定、任务轮询和 Socket.IO 播放器连接流程图",
        )
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(8)
        base.add_inline(caption, "图 1  云应用入口、鉴权与起流总流程", size=9, color=base.MUTED)
        return
    if lines and lines[0].strip() == "architectureDiagram":
        figure_path = Path(tempfile.gettempdir()) / "cmcc-cloudapp-reference-architecture.png"
        build_architecture_figure(figure_path)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(4)
        picture = paragraph.add_run().add_picture(str(figure_path), width=Inches(6.3))
        picture._inline.docPr.set(
            "descr",
            "浏览器、接入与编排平台、凭据和策略组件、10086、139 门户、云应用控制面及 Socket.IO WebRTC 的推荐分层架构图",
        )
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(0)
        caption.paragraph_format.space_after = Pt(8)
        base.add_inline(caption, "图 2  推荐分层架构与信任边界", size=9, color=base.MUTED)
        return
    base._original_add_code_block(document, lines)


def compact_final_notes(output_path: Path):
    document = Document(output_path)
    in_contents = False
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text == "目录":
            in_contents = True
            continue
        if in_contents:
            if paragraph._p.xpath('.//w:br[@w:type="page"]'):
                in_contents = False
            else:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(8.8)
        if paragraph.text.startswith(("来源文件：", "静态代码：", "说明：")):
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_together = True
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
        if paragraph.text.startswith("9. 超时、重试与异常处理"):
            # Keep the following multi-row table inside a fresh page budget.
            # LibreOffice otherwise clips its continuation rows when it starts
            # near the bottom of the preceding page.
            paragraph.paragraph_format.page_break_before = True
        if paragraph.style.name == "Code Block" and len(paragraph.text.splitlines()) <= 50:
            paragraph.paragraph_format.keep_together = True
        if index + 1 < len(paragraphs) and paragraphs[index + 1].style.name == "Code Block":
            paragraph.paragraph_format.keep_with_next = True
    document.save(output_path)


def main(output_dir: Path):
    source_dir = Path(__file__).resolve().parent
    output_dir.mkdir(parents=True, exist_ok=True)
    base.CJK_FONT = "Arial Unicode MS"
    base.add_cover = add_memo_masthead
    base._original_add_code_block = base.add_code_block
    base.add_code_block = add_code_or_flow
    output_path = output_dir / OUTPUT_NAME
    base.build(source_dir / SOURCE_NAME, output_path, METADATA)
    compact_final_notes(output_path)
    embed_font(output_path, EMBEDDED_CJK_FONT_PATH, base.CJK_FONT)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: _build_cloudapp_document.py <output-dir>")
    main(Path(sys.argv[1]))
