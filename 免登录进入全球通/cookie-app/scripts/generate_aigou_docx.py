#!/usr/bin/env python3
"""Build the three Aigou design documents from their Markdown sources."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"

DOCUMENTS = [
    {
        "source": "aigou-business-analysis.md",
        "output": "爱购商城业务订购分析文档.docx",
        "kind": "业务分析",
        "subtitle": "10086 统一认证至爱购权益商城直订链路",
        "audience": "产品、研发、测试、运营",
        "status": "已实现并通过自动化验证",
    },
    {
        "source": "aigou-order-api-design.md",
        "output": "爱购商城订购接口设计文档.docx",
        "kind": "接口设计",
        "subtitle": "cookie-app 爱购订单编排与上游接口规范",
        "audience": "后端、前端、联调、测试",
        "status": "基础路径 /api/aigou-orders",
    },
    {
        "source": "aigou-order-product-design.md",
        "output": "爱购商城业务订购产品设计文档.docx",
        "kind": "产品设计",
        "subtitle": "V2 业务订购中的爱购商城操作流程",
        "audience": "产品、设计、前端、验收",
        "status": "桌面端与移动端响应式交付",
    },
]

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "17212B"
MUTED = "5F6B76"
LIGHT_FILL = "F4F6F9"
CODE_FILL = "F2F4F7"
RULE = "AEB8C2"
BODY_FONT = "Heiti SC"
CJK_FONT = "Heiti SC"
MONO_FONT = "SFMono-Regular"


def set_run_font(run, *, name=BODY_FONT, east_asia=CJK_FONT, size=None,
                 color=None, bold=None, italic=None):
    run.font.name = name
    if run._element.get_or_add_rPr().rFonts is None:
        run._element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), east_asia)
    lang = run._element.get_or_add_rPr().find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        run._element.get_or_add_rPr().append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_border(paragraph, color=RULE, size="10"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, fallback, end])
    set_run_font(run, size=9, color=MUTED)


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal._element.rPr.append(language_element())
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25
    pf.widow_control = True

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style._element.rPr.append(language_element())
        pf = style.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.0
        pf.keep_with_next = True
        pf.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(INK)
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style._element.rPr.append(language_element())
        pf = style.paragraph_format
        pf.left_indent = Inches(0.375)
        pf.first_line_indent = Inches(-0.188)
        pf.space_before = Pt(0)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.25

    code = styles.add_style("Aigou Code", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = MONO_FONT
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string("26323D")
    code._element.rPr.rFonts.set(qn("w:ascii"), MONO_FONT)
    code._element.rPr.rFonts.set(qn("w:hAnsi"), MONO_FONT)
    code._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    code._element.rPr.append(language_element())
    code.paragraph_format.left_indent = Inches(0.12)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.08

    lead = styles.add_style("Aigou Lead", WD_STYLE_TYPE.PARAGRAPH)
    lead.font.name = BODY_FONT
    lead.font.size = Pt(11)
    lead.font.bold = True
    lead.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    lead._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT)
    lead._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT)
    lead._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    lead._element.rPr.append(language_element())
    lead.paragraph_format.space_before = Pt(0)
    lead.paragraph_format.space_after = Pt(10)
    lead.paragraph_format.line_spacing = 1.25


def language_element():
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")
    return lang


def configure_page(doc, kind):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    left = hp.add_run("爱购商城订购能力")
    set_run_font(left, size=9, color=MUTED, bold=True)
    right = hp.add_run(f"  |  {kind}")
    set_run_font(right, size=9, color=MUTED)
    set_paragraph_border(hp, color="D6DCE2", size="4")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    label = fp.add_run("内部参考  ·  ")
    set_run_font(label, size=9, color=MUTED)
    add_field(fp, "PAGE")
    sep = fp.add_run(" / ")
    set_run_font(sep, size=9, color=MUTED)
    add_field(fp, "NUMPAGES")


def add_inline_runs(paragraph, text, *, bold_default=False):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name=MONO_FONT, east_asia=CJK_FONT, size=9.2,
                         color=DARK_BLUE, bold=False)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EDF1F5")
            run._element.get_or_add_rPr().append(shading)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=11, color=INK, bold=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=11, color=INK, bold=bold_default)


def add_masthead(doc, config, title, version_line):
    version_line = version_line.replace("`", "")
    version_line = re.sub(r"^版本：\s*", "", version_line)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(7)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    kr = kicker.add_run(config["kind"].upper())
    set_run_font(kr, size=9, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(title)
    set_run_font(run, size=23, color="000000", bold=True)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(14)
    sr = sub.add_run(config["subtitle"])
    set_run_font(sr, size=12, color=MUTED)

    metadata = [
        ("版本", version_line),
        ("读者", config["audience"]),
        ("状态", config["status"]),
        ("证据", "HAR、页面包、cookie-app 源码与自动化测试"),
    ]
    for label, value in metadata:
        mp = doc.add_paragraph()
        mp.paragraph_format.space_after = Pt(2)
        mp.paragraph_format.line_spacing = 1.08
        lr = mp.add_run(f"{label}：")
        set_run_font(lr, size=10.5, color=INK, bold=True)
        vr = mp.add_run(value)
        set_run_font(vr, size=10.5, color=INK)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(6)
    rule.paragraph_format.space_after = Pt(9)
    set_paragraph_border(rule, color=BLUE, size="14")


def add_code_block(doc, lines):
    p = doc.add_paragraph(style="Aigou Code")
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), CODE_FILL)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9DEE3")
        border.set(qn("w:space"), "4")
        borders.append(border)
    p_pr.append(borders)
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, name=MONO_FONT, east_asia=CJK_FONT, size=8.5,
                     color="26323D")
        if idx < len(lines) - 1:
            run.add_break(WD_BREAK.LINE)


def parse_markdown(source_text):
    lines = source_text.splitlines()
    title = lines[0].removeprefix("# ").strip()
    version_line = lines[2].strip() if len(lines) > 2 else "1.0"
    body = lines[3:]
    return title, version_line, body


def add_markdown_body(doc, lines):
    paragraph_buffer = []
    code_lines = []
    in_code = False
    first_body_para = True

    def flush_paragraph():
        nonlocal paragraph_buffer, first_body_para
        if not paragraph_buffer:
            return
        text = " ".join(part.strip() for part in paragraph_buffer).strip()
        if text:
            style = "Aigou Lead" if first_body_para else "Normal"
            p = doc.add_paragraph(style=style)
            add_inline_runs(p, text, bold_default=first_body_para)
            if first_body_para:
                p_pr = p._p.get_or_add_pPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), LIGHT_FILL)
                p_pr.append(shading)
                p.paragraph_format.left_indent = Inches(0.12)
                p.paragraph_format.right_indent = Inches(0.12)
                first_body_para = False
        paragraph_buffer = []

    for raw in lines:
        line = raw.rstrip()
        if line.startswith("```"):
            flush_paragraph()
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line.strip():
            flush_paragraph()
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", line)
        if heading:
            flush_paragraph()
            level = min(len(heading.group(1)) - 1, 3)
            doc.add_paragraph(heading.group(2).strip(), style=f"Heading {level}")
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, bullet.group(1))
            continue
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if numbered:
            flush_paragraph()
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, numbered.group(1))
            continue
        paragraph_buffer.append(line)

    flush_paragraph()
    if in_code and code_lines:
        add_code_block(doc, code_lines)


def build(config):
    source = DOCS / config["source"]
    output = DOCS / config["output"]
    title, version_line, body = parse_markdown(source.read_text(encoding="utf-8"))

    doc = Document()
    configure_styles(doc)
    configure_page(doc, config["kind"])
    doc.core_properties.title = title
    doc.core_properties.subject = config["subtitle"]
    doc.core_properties.author = "cookie-app 项目组"
    doc.core_properties.keywords = "爱购商城, 10086, 业务订购, cookie-app"
    doc.core_properties.comments = "不含 Cookie、token、手机号、验证码或动态风控密文。"

    add_masthead(doc, config, title, version_line)
    add_markdown_body(doc, body)
    doc.save(output)
    return output


def main():
    outputs = [build(config) for config in DOCUMENTS]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    sys.exit(main())
